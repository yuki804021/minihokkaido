#!/usr/bin/env python3
"""北海道の路線・駅・運行パターンのデータ (data/network.js) を作る。

使い方:
    python3 tools/build_network.py

駅の並びは 駅データ.jp (piuccio/open-data-jp-railway-stations 経由)、
駅の位置と線路の形は「国土数値情報（鉄道データ N02、2024 年度）」（国土交通省）を使う。
取得したファイルは tools/cache/ に保存し、2 回目以降はそれを使う。

運転間隔・停車駅・フェリーや航空便の時刻は、すべて制作者による推計。
"""
import heapq
import json
import math
import os
import re
import unicodedata
import urllib.request
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
CACHE = os.path.join(HERE, 'cache')
STATIONS_URL = 'https://raw.githubusercontent.com/piuccio/open-data-jp-railway-stations/master/stations.json'
N02_URL = 'https://nlftp.mlit.go.jp/ksj/gml/data/N02/N02-24/N02-24_GML.zip'

R = 6371008.8


# ------------------------------------------------------------------ 幾何
def haversine(a, b):
    la1, la2 = math.radians(a[1]), math.radians(b[1])
    dla, dlo = la2 - la1, math.radians(b[0] - a[0])
    h = math.sin(dla / 2) ** 2 + math.cos(la1) * math.cos(la2) * math.sin(dlo / 2) ** 2
    return 2 * R * math.asin(math.sqrt(h))


def bearing(a, b):
    la1, la2 = math.radians(a[1]), math.radians(b[1])
    dlo = math.radians(b[0] - a[0])
    y = math.sin(dlo) * math.cos(la2)
    x = math.cos(la1) * math.sin(la2) - math.sin(la1) * math.cos(la2) * math.cos(dlo)
    return math.degrees(math.atan2(y, x)) % 360


def offset(c, brg_deg, dist):
    b = math.radians(brg_deg)
    dlat = dist * math.cos(b) / R
    dlon = dist * math.sin(b) / (R * math.cos(math.radians(c[1])))
    return [c[0] + math.degrees(dlon), c[1] + math.degrees(dlat)]


def rnd(c):
    return [round(c[0], 6), round(c[1], 6)]


def path_length(coords):
    return sum(haversine(coords[i], coords[i + 1]) for i in range(len(coords) - 1))


def densify(coords, step):
    """折れ線を step [m] ごとの点に分ける (航路・飛行経路用)"""
    out = [coords[0]]
    for a, b in zip(coords, coords[1:]):
        n = max(1, int(haversine(a, b) // step))
        for k in range(1, n):
            f = k / n
            out.append([a[0] + (b[0] - a[0]) * f, a[1] + (b[1] - a[1]) * f])
        out.append(list(b))
    return out


# ------------------------------------------------------------------ 駅データ
def load_stations():
    os.makedirs(CACHE, exist_ok=True)
    path = os.path.join(CACHE, 'stations.json')
    if not os.path.exists(path):
        print('download', STATIONS_URL)
        urllib.request.urlretrieve(STATIONS_URL, path)
    with open(path, encoding='utf-8') as f:
        groups = json.load(f)
    lines = {}
    for g in groups:
        for s in g['stations']:
            lines.setdefault(s['ekidata_line_id'], []).append(s)
    for v in lines.values():
        v.sort(key=lambda s: int(s['ekidata_id']))
    return {k: [(s['name_kanji'], [round(s['lon'], 6), round(s['lat'], 6)]) for s in v] for k, v in lines.items()}


EKI = load_stations()

# 駅データ.jp の収録後に改称された駅
RENAMED = {'石狩太美': '太美', '石狩当別': '当別', '東風連': '名寄高校'}
# 駅データ.jp に無い駅 (位置は国土数値情報から取る)
EXTRA = {
    '奥津軽いまべつ': [140.5153, 41.1453],
}
# 駅データ.jp に無く、国土数値情報にある駅: 路線ごとに、いちばん近い区間へ挿入する
ADDED = {'jr_sassho': ['ロイズタウン']}
# 廃止された駅は、国土数値情報 (2024 年度) に無い JR の駅として自動で除く (refine_stations)。
# 廃止された区間は、下の LINES で使う区間を選ぶことで除いている


def station_list(line_id):
    return [(RENAMED.get(n, n), c) for n, c in EKI[line_id]]


def take(line_id, a, b):
    """路線 line_id の a 駅から b 駅まで (逆順も可)"""
    lst = station_list(line_id)
    names = [n for n, _ in lst]
    i, j = names.index(a), names.index(b)
    return lst[i:j + 1] if i <= j else list(reversed(lst[j:i + 1]))


def pick(line_id, *names):
    """路線 line_id から駅を名前の順に取り出す (並びを直すとき用)"""
    d = dict(station_list(line_id))
    return [(n, d[n]) for n in names]


def join(*parts):
    out = []
    for p in parts:
        for s in p:
            if out and out[-1][0] == s[0]:
                continue
            out.append(s)
    return out


# ------------------------------------------------------------------ 国土数値情報 (鉄道データ)
def norm(name):
    return unicodedata.normalize('NFKC', name).replace('ヶ', 'ケ').replace('ヵ', 'カ')


def midpoint(coords):
    half, acc = path_length(coords) / 2, 0.0
    for a, b in zip(coords, coords[1:]):
        d = haversine(a, b)
        if acc + d >= half and d > 0:
            f = (half - acc) / d
            return [a[0] + (b[0] - a[0]) * f, a[1] + (b[1] - a[1]) * f]
        acc += d
    return coords[0]


def load_n02():
    files = {k: os.path.join(CACHE, f'N02-24_{k}.geojson') for k in ('Station', 'RailroadSection')}
    if not all(os.path.exists(p) for p in files.values()):
        zpath = os.path.join(CACHE, 'N02-24_GML.zip')
        if not os.path.exists(zpath):
            print('download', N02_URL)
            urllib.request.urlretrieve(N02_URL, zpath)
        with zipfile.ZipFile(zpath) as z:
            for k, p in files.items():
                with z.open(f'UTF-8/N02-24_{k}.geojson') as src, open(p, 'wb') as dst:
                    dst.write(src.read())

    def near(coords):  # 北海道と青森県北部だけを使う
        return any(139 < x < 146.5 and 40.5 < y < 46 for x, y in coords)

    with open(files['Station'], encoding='utf-8') as f:
        st = [(norm(x['properties']['N02_005']), x['properties']['N02_004'], x['properties']['N02_003'],
               rnd(midpoint(x['geometry']['coordinates'])))
              for x in json.load(f)['features'] if near(x['geometry']['coordinates'])]
    with open(files['RailroadSection'], encoding='utf-8') as f:
        rs = [(x['properties']['N02_004'], x['properties']['N02_003'], x['geometry']['coordinates'])
              for x in json.load(f)['features'] if near(x['geometry']['coordinates'])]
    return st, rs


N02_STATIONS, N02_SECTIONS = load_n02()

JR = {'北海道旅客鉄道'}
# 路線ごとに使う国土数値情報の事業者と路線名 (路線名 None はその事業者の全路線)。
# 路線名の先頭はその路線自身の線路で、駅はまずこの線路に吸着させる
N02_OF = {
    'subway_namboku': ({'札幌市'}, ['南北線']),
    'subway_tozai': ({'札幌市'}, ['東西線']),
    'subway_toho': ({'札幌市'}, ['東豊線']),
    'sapporo_tram': ({'一般社団法人札幌市交通事業振興公社'}, None),
    'hakodate_tram2': ({'函館市'}, None),
    'hakodate_tram5': ({'函館市'}, None),
    'jr_hakodate_s': (JR, ['函館線']),
    'jr_hakodate_sawara': (JR, ['函館線']),
    'jr_hakodate_yama': (JR, ['函館線']),
    'jr_hakodate_n': (JR, ['函館線']),
    'jr_muroran': (JR, ['室蘭線']),
    'jr_muroran_branch': (JR, ['室蘭線']),
    'jr_muroran_n': (JR, ['室蘭線', '千歳線']),
    'jr_chitose': (JR, ['千歳線', '函館線', '室蘭線']),
    'jr_chitose_airport': (JR, ['千歳線']),
    'jr_sekisho': (JR, ['石勝線', '千歳線', '根室線']),
    'jr_nemuro_w': (JR, ['根室線', '函館線']),
    'jr_nemuro': (JR, ['根室線', '石勝線']),
    'jr_hanasaki': (JR, ['根室線']),
    'jr_hidaka': (JR, ['日高線', '室蘭線']),
    'jr_sassho': (JR, ['札沼線', '函館線']),
    'jr_furano': (JR, ['富良野線', '函館線', '根室線']),
    'jr_soya': (JR, ['宗谷線', '函館線']),
    'jr_sekihoku': (JR, ['石北線', '宗谷線', '函館線']),
    'jr_senmo': (JR, ['釧網線', '根室線', '石北線']),
    'shinkansen': ({'北海道旅客鉄道', '東日本旅客鉄道'}, ['北海道新幹線', '海峡線', '東北新幹線']),
    'isaribi': ({'道南いさりび鉄道', '北海道旅客鉄道'}, ['道南いさりび鉄道線', '函館線']),
}


def refine_stations(line_id, stations):
    """駅の位置を国土数値情報の駅 (ホームの中点) に合わせ、国土数値情報に無い JR の駅 (廃止駅) を除く"""
    ops, names = N02_OF[line_id]
    out = []
    for n, c in stations:
        # 同じ名前の駅が別の路線にもある (例: 地下鉄 さっぽろ) ので、その路線自身の駅を優先する
        cands = [(names is not None and x[2] != names[0], haversine(c, x[3]), x[3]) for x in N02_STATIONS
                 if x[0] == norm(n) and x[1] in ops and (names is None or x[2] in names)]
        cands = [z for z in cands if z[1] < 3000]
        if cands:
            out.append((n, min(cands)[2]))
        elif n in EXTRA or not ops & JR:
            out.append((n, c))
            if n not in EXTRA:
                print(f'  {line_id}: {n} は国土数値情報に無いため、駅データ.jp の位置を使う')
        else:
            print(f'  {line_id}: {n} は国土数値情報 (2024 年度) に無いため除外')
    for n in ADDED.get(line_id, []):
        c = min((x for x in N02_STATIONS if x[0] == norm(n) and x[1] in ops),
                key=lambda x: haversine(x[3], out[0][1]))[3]
        # 挿入して増える距離がいちばん小さい区間に入れる
        k = min(range(1, len(out)), key=lambda i: haversine(out[i - 1][1], c) + haversine(c, out[i][1]) - haversine(out[i - 1][1], out[i][1]))
        out.insert(k, (n, c))
    return out


class TrackGraph:
    """国土数値情報の線路 (LineString) をつないだグラフ。駅を最寄りの線路に吸着させ、駅間を最短経路で結ぶ"""
    SNAP = 350   # 駅から線路までの距離の上限 [m]
    LINK = 30    # 線路の端どうしがこれより近ければ、つながっているとみなす [m]
    CELL = 0.02  # 空間索引のマス [度]

    def __init__(self, ops, names):
        self.pts, self.key = [], {}
        self.segs = []  # (点 a, 点 b)
        self.own = []   # その路線自身の線路か
        for op, line, coords in N02_SECTIONS:
            if op not in ops or (names is not None and line not in names):
                continue
            ids = [self._node(c) for c in coords]
            for a, b in zip(ids, ids[1:]):
                if a != b:
                    self.segs.append((a, b))
                    self.own.append(names is None or line == names[0])
        self.grid = {}
        for i, (a, b) in enumerate(self.segs):
            for cell in self._cells(self.pts[a], self.pts[b]):
                self.grid.setdefault(cell, []).append(i)
        self.links = []
        ends = {}
        for a, b in self.segs:
            for p in (a, b):
                ends.setdefault(self._cell(self.pts[p]), set()).add(p)
        for cell, ps in ends.items():
            near = set().union(*(ends.get((cell[0] + dx, cell[1] + dy), set()) for dx in (-1, 0, 1) for dy in (-1, 0, 1)))
            for p in ps:
                for q in near:
                    if p < q and haversine(self.pts[p], self.pts[q]) < self.LINK:
                        self.links.append((p, q))
        self.snaps = {}  # 線分 → [(t, 点)]

    def _node(self, c):
        k = (round(c[0], 6), round(c[1], 6))
        if k not in self.key:
            self.key[k] = len(self.pts)
            self.pts.append(list(k))
        return self.key[k]

    def _cell(self, c):
        return (int(c[0] // self.CELL), int(c[1] // self.CELL))

    def _cells(self, a, b):
        x0, y0 = self._cell([min(a[0], b[0]), min(a[1], b[1])])
        x1, y1 = self._cell([max(a[0], b[0]), max(a[1], b[1])])
        return [(x, y) for x in range(x0, x1 + 1) for y in range(y0, y1 + 1)]

    def snap(self, c):
        """c から最も近い線路上の点を、グラフの点として登録する。遠すぎれば None"""
        cx, cy = self._cell(c)
        kx = 111320 * math.cos(math.radians(c[1]))
        best = None
        for dx in (-1, 0, 1):
            for dy in (-1, 0, 1):
                for i in self.grid.get((cx + dx, cy + dy), []):
                    a, b = self.pts[self.segs[i][0]], self.pts[self.segs[i][1]]
                    ax, ay = (a[0] - c[0]) * kx, (a[1] - c[1]) * 110540
                    bx, by = (b[0] - c[0]) * kx, (b[1] - c[1]) * 110540
                    vx, vy = bx - ax, by - ay
                    L2 = vx * vx + vy * vy
                    t = 0.0 if L2 == 0 else max(0.0, min(1.0, -(ax * vx + ay * vy) / L2))
                    d = math.hypot(ax + vx * t, ay + vy * t)
                    # その路線自身の線路を優先する (乗換駅で隣の路線に吸着しないように)
                    rank = (not (self.own[i] and d <= self.SNAP), d)
                    if best is None or rank < best[3]:
                        best = (d, i, t, rank)
        if best is None or best[0] > self.SNAP:
            return None
        _, i, t, _ = best
        a, b = self.segs[i]
        if t <= 0:
            return a
        if t >= 1:
            return b
        pa, pb = self.pts[a], self.pts[b]
        p = self._node([pa[0] + (pb[0] - pa[0]) * t, pa[1] + (pb[1] - pa[1]) * t])
        self.snaps.setdefault(i, []).append((t, p))
        return p

    def build(self):
        """吸着させた点で線分を分けて、隣接リストを作る"""
        self.adj = [[] for _ in self.pts]

        def edge(p, q):
            d = haversine(self.pts[p], self.pts[q])
            self.adj[p].append((q, d))
            self.adj[q].append((p, d))
        for i, (a, b) in enumerate(self.segs):
            chain = [a] + [p for _, p in sorted(self.snaps.get(i, []))] + [b]
            for p, q in zip(chain, chain[1:]):
                if p != q:
                    edge(p, q)
        for p, q in self.links:
            edge(p, q)

    def route(self, s, t):
        dist, prev = {s: 0.0}, {}
        heap = [(0.0, s)]
        while heap:
            d, u = heapq.heappop(heap)
            if u == t:
                break
            if d > dist[u]:
                continue
            for v, w in self.adj[u]:
                nd = d + w
                if nd < dist.get(v, float('inf')):
                    dist[v], prev[v] = nd, u
                    heapq.heappush(heap, (nd, v))
        if t not in dist:
            return None
        path = [t]
        while path[-1] != s:
            path.append(prev[path[-1]])
        return [self.pts[p] for p in reversed(path)]


def track_shapes(line_id, stations):
    """駅と駅の間の線路の形 (途中の点の並び) を、区間ごとに返す"""
    g = TrackGraph(*N02_OF[line_id])
    nodes = [g.snap(c) for _, c in stations]
    g.build()
    shapes = []
    for k in range(len(stations) - 1):
        (na, ca), (nb, cb) = stations[k], stations[k + 1]
        straight = haversine(ca, cb)
        via = None
        if nodes[k] is not None and nodes[k + 1] is not None:
            via = g.route(nodes[k], nodes[k + 1])
            if via is None:
                print(f'  {line_id}: {na}〜{nb} は線路がつながっていないため直線にする')
        else:
            print(f'  {line_id}: {na}〜{nb} は線路に吸着できないため直線にする')
        if via is not None:
            ratio = path_length([ca] + via + [cb]) / max(straight, 1)
            if ratio > 2.5 and straight > 500:
                print(f'  {line_id}: {na}〜{nb} の経路が遠回り ({ratio:.1f} 倍) のため直線にする')
                via = None
        if via is None:
            shapes.append([])
            continue
        # 端の点が駅とほぼ同じ位置なら省く
        via = [rnd(p) for p in via]
        if via and haversine(via[0], ca) < 5:
            via = via[1:]
        if via and haversine(via[-1], cb) < 5:
            via = via[:-1]
        shapes.append(via)
    return shapes


# ------------------------------------------------------------------ 路線 (線路)
GROUPS = [
    {'id': 'sapporo_subway', 'name': '札幌市営地下鉄'},
    {'id': 'tram', 'name': '路面電車（札幌・函館）'},
    {'id': 'jr', 'name': 'JR北海道 普通・快速'},
    {'id': 'jr_ltd', 'name': 'JR北海道 特急・北海道新幹線'},
    {'id': 'hokkaido_other', 'name': '道南いさりび鉄道'},
    {'id': 'highway_bus', 'name': '都市間高速バス'},
    # 路線バスは動かさず、路線図・停留所・時刻表のみ (data/bus_map.js、tools/build_bus.py で作る)
    {'id': 'route_bus', 'name': '路線バス（路線図・停留所）', 'color': '#8D6E63'},
    {'id': 'ferry', 'name': 'フェリー・旅客船'},
    {'id': 'air', 'name': '航空便（推計）'},
]

JR_COLOR = '#43A047'

hakodate_main = join(take('11101', '函館', '駒ケ岳'), pick('11101', '森'), take('11101', '森', '長万部'))
hakodate_sawara = join(pick('11101', '大沼'), take('11101', '池田園', '東森'), pick('11101', '森'))
muroran_main = join(take('11104', '長万部', '本輪西'), take('11104', '東室蘭', '苫小牧'))
muroran_branch = pick('11104', '東室蘭', '輪西', '御崎', '母恋', '室蘭')
chitose_main = join(pick('11109', '沼ノ端', '植苗'), take('11109', '南千歳', '白石'), pick('11109', '苗穂', '札幌'))
chitose_airport = pick('11109', '南千歳', '新千歳空港')
sekisho = join(pick('11110', '南千歳', '追分', '川端', '滝ノ上', '新夕張', '占冠', 'トマム'), pick('11107', '新得'))
shinkansen = [('新函館北斗', dict(EKI['11101'])['新函館北斗']), ('木古内', dict(EKI['99108'])['木古内']),
              ('奥津軽いまべつ', EXTRA['奥津軽いまべつ']), ('新青森', dict(EKI['11202'])['新青森'])]
sapporo_tram = station_list('99104')
sapporo_tram = sapporo_tram + [sapporo_tram[0]]  # 環状運転 (2015 年にループ化)

LINES = [
    # id, 名前, 事業者, グループ, 種別, 色, 駅
    ('subway_namboku', '札幌市営地下鉄 南北線', '札幌市交通局', 'sapporo_subway', 'rail', '#2E9E48', station_list('99102')),
    ('subway_tozai', '札幌市営地下鉄 東西線', '札幌市交通局', 'sapporo_subway', 'rail', '#F39800', station_list('99101')),
    ('subway_toho', '札幌市営地下鉄 東豊線', '札幌市交通局', 'sapporo_subway', 'rail', '#0091D5', station_list('99103')),
    ('sapporo_tram', '札幌市電', '札幌市交通局', 'tram', 'tram', '#7CB342', sapporo_tram),
    ('hakodate_tram2', '函館市電 2系統', '函館市企業局', 'tram', 'tram', '#E57373', station_list('99105')),
    ('hakodate_tram5', '函館市電 5系統', '函館市企業局', 'tram', 'tram', '#4FC3F7', station_list('99106')),
    ('jr_hakodate_s', 'JR函館本線 (函館〜長万部)', 'JR北海道', 'jr', 'rail', JR_COLOR, hakodate_main),
    ('jr_hakodate_sawara', 'JR函館本線 (砂原支線)', 'JR北海道', 'jr', 'rail', JR_COLOR, hakodate_sawara),
    ('jr_hakodate_yama', 'JR函館本線 (長万部〜小樽)', 'JR北海道', 'jr', 'rail', JR_COLOR, station_list('11102')),
    ('jr_hakodate_n', 'JR函館本線 (小樽〜旭川)', 'JR北海道', 'jr', 'rail', JR_COLOR, station_list('11103')),
    ('jr_muroran', 'JR室蘭本線 (長万部〜苫小牧)', 'JR北海道', 'jr', 'rail', JR_COLOR, muroran_main),
    ('jr_muroran_branch', 'JR室蘭本線 (東室蘭〜室蘭)', 'JR北海道', 'jr', 'rail', JR_COLOR, muroran_branch),
    ('jr_muroran_n', 'JR室蘭本線 (苫小牧〜岩見沢)', 'JR北海道', 'jr', 'rail', JR_COLOR, station_list('11105')),
    ('jr_chitose', 'JR千歳線', 'JR北海道', 'jr', 'rail', JR_COLOR, chitose_main),
    ('jr_chitose_airport', 'JR千歳線 (南千歳〜新千歳空港)', 'JR北海道', 'jr', 'rail', JR_COLOR, chitose_airport),
    ('jr_sekisho', 'JR石勝線', 'JR北海道', 'jr', 'rail', JR_COLOR, sekisho),
    ('jr_nemuro_w', 'JR根室本線 (滝川〜富良野)', 'JR北海道', 'jr', 'rail', JR_COLOR, take('11106', '滝川', '富良野')),
    ('jr_nemuro', 'JR根室本線 (新得〜釧路)', 'JR北海道', 'jr', 'rail', JR_COLOR, station_list('11107')),
    ('jr_hanasaki', 'JR根室本線 (花咲線)', 'JR北海道', 'jr', 'rail', JR_COLOR, station_list('11108')),
    ('jr_hidaka', 'JR日高本線', 'JR北海道', 'jr', 'rail', JR_COLOR, take('11111', '苫小牧', '鵡川')),
    ('jr_sassho', 'JR札沼線 (学園都市線)', 'JR北海道', 'jr', 'rail', JR_COLOR, take('11112', '札幌', '北海道医療大学')),
    ('jr_furano', 'JR富良野線', 'JR北海道', 'jr', 'rail', JR_COLOR, station_list('11114')),
    ('jr_soya', 'JR宗谷本線', 'JR北海道', 'jr', 'rail', JR_COLOR, station_list('11115')),
    ('jr_sekihoku', 'JR石北本線', 'JR北海道', 'jr', 'rail', JR_COLOR, station_list('11116')),
    ('jr_senmo', 'JR釧網本線', 'JR北海道', 'jr', 'rail', JR_COLOR, station_list('11117')),
    ('shinkansen', '北海道新幹線 (新函館北斗〜新青森)', 'JR北海道', 'jr_ltd', 'rail', '#7B5EA7', shinkansen),
    ('isaribi', '道南いさりび鉄道線', '道南いさりび鉄道', 'hokkaido_other', 'rail', '#1565C0', station_list('99108')),
]
print('国土数値情報で駅の位置と線路の形を合わせる')
LINES = [(*l[:6], refine_stations(l[0], l[6])) for l in LINES]
SHAPES = {l[0]: track_shapes(l[0], l[6]) for l in LINES}
LINE = {l[0]: l for l in LINES}


def stations_of(line_id):
    return LINE[line_id][6]


def seg(line_id, a, b):
    """路線の a 駅から b 駅まで。各駅に、直前の駅からの線路の形 (途中の点) を添えて返す"""
    lst = stations_of(line_id)
    shapes = SHAPES[line_id]
    names = [n for n, _ in lst]
    i, j = names.index(a), names.index(b)
    if i <= j:
        return [(lst[k][0], lst[k][1], shapes[k - 1] if k > i else []) for k in range(i, j + 1)]
    return [(lst[k][0], lst[k][1], list(reversed(shapes[k])) if k < i else []) for k in range(i, j - 1, -1)]


def rows_of(path, stops=None):
    """seg をつないだ経路を、sim.js の path ([名前, 座標, 停車]) にする。途中の点は名前なし"""
    rows = []
    for i, (n, c, via) in enumerate(path):
        rows.extend(['', p, 0] for p in via)
        stop = 1 if stops is None or n in stops or i in (0, len(path) - 1) else 0
        rows.append([n, c, stop])
    return rows


def shape_of(line_id):
    path = seg(line_id, stations_of(line_id)[0][0], stations_of(line_id)[-1][0]) if line_id != 'sapporo_tram' else \
        [(n, c, SHAPES[line_id][k - 1] if k else []) for k, (n, c) in enumerate(stations_of(line_id))]
    return [r[1] for r in rows_of(path)]


# ------------------------------------------------------------------ 運行系統
RAIL = {'kind': 'rail', 'cars': 2, 'carLength': 21, 'width': 2.9, 'height': 4.0, 'speed': 75, 'dwell': 30, 'accel': 30}
LOCAL1 = {**RAIL, 'cars': 1}
# 特急の最高速度は、主な区間の所要時間が実際に近くなるよう系統ごとに調整している (LIMITED の speed)
LTD = {**RAIL, 'cars': 5, 'speed': 100, 'dwell': 60, 'accel': 40}
SUBWAY = {'kind': 'rail', 'cars': 6, 'carLength': 18, 'width': 3.1, 'height': 3.7, 'speed': 50, 'dwell': 20, 'accel': 15}
TRAM = {'kind': 'tram', 'cars': 1, 'carLength': 13, 'width': 2.4, 'height': 3.8, 'speed': 16, 'dwell': 20, 'accel': 8}

services = []


def service(sid, name, group, line, color, legs, spec, stops=None, **extra):
    """legs: [(路線, 始点, 終点), ...] をつないだ経路。stops を省くと各駅に停車"""
    path = join(*[seg(*leg) for leg in legs])
    names = [s[0] for s in path]
    if stops is not None:
        missing = [s for s in stops if s not in names]
        assert not missing, (sid, missing)
    rows = rows_of(path, stops)
    sv = {'id': sid, 'name': name, 'group': group, 'line': line, 'color': color, **spec,
          'loop': False, 'both': True, 'offset': 0, 'path': rows}
    sv.update(extra)
    services.append(sv)
    return sv


# 札幌市営地下鉄 (運転間隔は推計)
service('namboku', '南北線', 'sapporo_subway', 'subway_namboku', '#2E9E48', [('subway_namboku', '麻生', '真駒内')],
        {**SUBWAY, 'cars': 6}, bands=[['06:00', '07:30', 6], ['07:30', '09:00', 4], ['09:00', '17:00', 7], ['17:00', '19:30', 5], ['19:30', '24:00', 8]])
service('tozai', '東西線', 'sapporo_subway', 'subway_tozai', '#F39800', [('subway_tozai', '宮の沢', '新さっぽろ')],
        {**SUBWAY, 'cars': 7}, bands=[['06:00', '07:30', 6], ['07:30', '09:00', 4], ['09:00', '17:00', 7], ['17:00', '19:30', 5], ['19:30', '24:00', 8]])
service('toho', '東豊線', 'sapporo_subway', 'subway_toho', '#0091D5', [('subway_toho', '栄町', '福住')],
        {**SUBWAY, 'cars': 4}, bands=[['06:00', '07:30', 8], ['07:30', '09:00', 5], ['09:00', '17:00', 8], ['17:00', '19:30', 6], ['19:30', '24:00', 10]])

# 札幌市電 (ループ線を外回り・内回りで運転)
loop = rows_of([(n, c, SHAPES['sapporo_tram'][k - 1] if k else []) for k, (n, c) in enumerate(stations_of('sapporo_tram'))])
for sid, name, path in (('sapporo_tram_out', '外回り', loop), ('sapporo_tram_in', '内回り', list(reversed(loop)))):
    services.append({'id': sid, 'name': f'札幌市電 {name}', 'group': 'tram', 'line': 'sapporo_tram', 'color': '#7CB342', **TRAM,
                     'loop': True, 'both': False, 'offset': 0 if sid.endswith('out') else 3, 'path': path,
                     'bands': [['06:30', '23:00', 8]]})

# 函館市電: 時刻表データ (GTFS) の便ごとの時刻で走らせる。取得できなければ推計ダイヤ
# ファイルの URL はダイヤ改正ごとに date が変わるので、ダウンロードするときにデータセットのページから最新のものを探す
HAKODATE_DATASET = 'https://ckan.odpt.org/dataset/hakodate_city_alllines'
HAKODATE_GTFS_URL = 'https://api-public.odpt.org/api/v4/files/odpt/HakodateCity/Alllines.zip?date=20260815'  # 見つからないとき
CREDITS = []


def latest_hakodate_url():
    import re

    def get(url):
        with urllib.request.urlopen(url, timeout=30) as r:
            return r.read().decode('utf-8', 'replace')
    try:
        found = set()
        for res in set(re.findall(r'/dataset/hakodate_city_alllines/resource/[0-9a-f-]+', get(HAKODATE_DATASET))):
            found |= set(re.findall(r'https://api-public\.odpt\.org/api/v4/files/odpt/HakodateCity/Alllines\.zip\?date=\d{8}',
                                    get('https://ckan.odpt.org' + res)))
        if found:
            return max(found, key=lambda u: u[-8:])
    except OSError as e:
        print('  函館市電のデータセットのページを読めないため、既定の URL を使う:', e)
    return HAKODATE_GTFS_URL

with open(os.path.join(HERE, 'holidays.json'), encoding='utf-8') as f:
    HOLIDAYS = set(json.load(f))


def read_gtfs(zpath):
    import csv
    import io
    with zipfile.ZipFile(zpath) as z:
        def table(name):
            member = next((x for x in z.namelist() if x.split('/')[-1] == name), None)
            if member is None:
                return []
            with z.open(member) as f:
                return list(csv.DictReader(io.TextIOWrapper(f, encoding='utf-8-sig')))
        return {k: table(k + '.txt') for k in ('stops', 'routes', 'trips', 'stop_times', 'calendar', 'calendar_dates')}


def gtfs_sec(s):
    h, m, sec = map(int, s.strip().split(':'))
    return h * 3600 + m * 60 + sec


def day_type_of(ymd):
    import datetime
    d = datetime.date(int(ymd[:4]), int(ymd[4:6]), int(ymd[6:]))
    return 'holiday' if d.weekday() >= 5 or ymd in HOLIDAYS else 'weekday'


def service_days(g):
    """service_id → その便が走るダイヤの種類 ('weekday' / 'holiday')"""
    days = {}
    for c in g['calendar']:
        ds = set()
        if any(c.get(d) == '1' for d in ('monday', 'tuesday', 'wednesday', 'thursday', 'friday')):
            ds.add('weekday')
        if any(c.get(d) == '1' for d in ('saturday', 'sunday')):
            ds.add('holiday')
        days[c['service_id']] = ds
    # calendar.txt に無い service_id は、calendar_dates.txt の運行日から判定する
    for c in g['calendar_dates']:
        if c.get('exception_type') == '1' and c['service_id'] not in {x['service_id'] for x in g['calendar']}:
            days.setdefault(c['service_id'], set()).add(day_type_of(c['date']))
    return {k: sorted(v) for k, v in days.items()}


def gtfs_services(zpath, line_ids, prefix, group, spec, colors, alias=None, route_name=None):
    """GTFS の便を、停車駅の並びが同じものごとに 1 つの系統にまとめる。経路は line_ids の路線の線路を使う。
    alias: GTFS の停留場名 → 路線データの駅名 (略称などの読み替え)"""
    g = read_gtfs(zpath)
    alias = {norm(k): norm(v) for k, v in (alias or {}).items()}
    stop_name = {s['stop_id']: alias.get(norm(s['stop_name']), norm(s['stop_name'])) for s in g['stops']}
    route_of = {r['route_id']: r for r in g['routes']}
    days = service_days(g)
    times = {}
    for r in g['stop_times']:
        times.setdefault(r['trip_id'], []).append(r)
    patterns = {}
    for t in g['trips']:
        st = sorted(times.get(t['trip_id'], []), key=lambda r: int(r['stop_sequence']))
        if len(st) < 2:
            continue
        names = tuple(stop_name[r['stop_id']] for r in st)
        patterns.setdefault((t['route_id'], names), []).append((t, st))
    out, unmatched = [], []
    for k, ((route_id, names), trips) in enumerate(sorted(patterns.items(), key=lambda x: (x[0][0], x[0][1]))):
        found = None
        for lid in line_ids:
            lst = [norm(n) for n, _ in stations_of(lid)]
            if all(n in lst for n in names):
                idx = [lst.index(n) for n in names]
                if idx == sorted(idx) or idx == sorted(idx, reverse=True):
                    found = (lid, idx)
                    break
        if found is None:
            unmatched.append((names, len(trips)))
            continue
        lid, idx = found
        st_names = [stations_of(lid)[i][0] for i in idx]
        path = seg(lid, st_names[0], st_names[-1])
        rows = rows_of(path, set(st_names))
        trip_rows = []
        for t, st in trips:
            dep0 = gtfs_sec(st[0]['departure_time'] or st[0]['arrival_time'])
            tt = [[gtfs_sec(r['arrival_time'] or r['departure_time']) - dep0, gtfs_sec(r['departure_time'] or r['arrival_time']) - dep0] for r in st]
            trip_rows.append({'dep': dep0, 't': tt, 'days': days.get(t['service_id'], ['weekday', 'holiday'])})
        trip_rows.sort(key=lambda x: x['dep'])
        route = route_of.get(route_id, {})
        rname = route.get('route_short_name') or route.get('route_long_name') or ''
        if route_name:
            rname = route_name(rname)
        color = colors(rname, route)
        out.append({**spec, 'id': f'{prefix}{k}', 'name': f'{rname} {st_names[0]}〜{st_names[-1]}'.strip(), 'group': group,
                    'line': lid, 'color': color, 'loop': False, 'both': False, 'offset': 0, 'path': rows, 'trips': trip_rows})
    for names, n in unmatched:
        print(f'  {prefix}: 路線の駅と合わない便 {n} 本を除外: {" → ".join(names)}')
    return out


def hakodate_color(rname, route):
    if route.get('route_color'):
        return '#' + route['route_color'].lstrip('#')
    return '#E57373' if '2' in rname else '#4FC3F7'  # route_color が無いとき


hakodate_zip = os.path.join(CACHE, 'hakodate_tram_gtfs.zip')
try:
    if not os.path.exists(hakodate_zip):
        url = latest_hakodate_url()
        print('download', url)
        urllib.request.urlretrieve(url, hakodate_zip)
    # GTFS の系統名は ② ⑤、停留場名は「アリーナ前」(正式には 函館アリーナ前)
    hakodate = gtfs_services(hakodate_zip, ['hakodate_tram2', 'hakodate_tram5'], 'hakodate_gtfs', 'tram', TRAM, hakodate_color,
                             alias={'アリーナ前': '函館アリーナ前'},
                             route_name=lambda r: f'{unicodedata.normalize("NFKC", r)}系統' if r in '①②③④⑤' else r)
except OSError as e:
    print('  函館市電の時刻表データを取得できないため、推計ダイヤにする:', e)
    hakodate = []
if hakodate:
    services.extend(hakodate)
    CREDITS.append('函館市企業局 函館市電 GTFS')
else:
    service('hakodate2', '2系統 湯の川〜谷地頭', 'tram', 'hakodate_tram2', '#E57373', [('hakodate_tram2', '湯の川', '谷地頭')],
            TRAM, bands=[['06:30', '22:30', 16]])
    service('hakodate5', '5系統 湯の川〜函館どつく前', 'tram', 'hakodate_tram5', '#4FC3F7', [('hakodate_tram5', '湯の川', '函館どつく前')],
            TRAM, bands=[['06:30', '22:30', 12]], offset=6)

# ------------------------------------------------------------------ 都市間高速バス
# 北海道オープンデータプラットフォーム (HODA) の「高速バス」と、根室交通の札幌根室線の時刻表データ (GTFS) で走らせる。
# 経路は、停留所を OpenStreetMap の道路 (OSRM) に沿ってつないだもの。路線バスは build_bus.py (地図と時刻表のみ)
import sys  # noqa: E402
sys.path.insert(0, HERE)
import gtfs_util as gu  # noqa: E402

HIGHWAY_LINES = []
HIGHWAY_COLOR = '#E65100'
BUS = {'kind': 'bus', 'cars': 1, 'carLength': 12, 'width': 2.5, 'height': 3.4, 'speed': 70, 'dwell': 20, 'accel': 20}
NEMURO_GTFS_URL = 'https://api.gtfs-data.jp/v2/organizations/nemurokotsu/feeds/nemurobus/files/feed.zip'
NEMURO_HIGHWAY_ROUTES = {'札幌線_R'}  # 根室交通のうち都市間高速バス (札幌根室線)


def highway_services(zpath, prefix, keep_route=None):
    g = gu.read_gtfs(zpath)
    ref = gu.reference_dates()
    run = gu.runs_on(g)
    stops = {s['stop_id']: s for s in g['stops']}
    routes = {r['route_id']: r for r in g['routes']}
    agency = {a['agency_id']: a['agency_name'] for a in g['agency']}
    times = gu.trips_by_id(g)
    patterns = {}
    for t in g['trips']:
        route = routes[t['route_id']]
        if keep_route and not keep_route(route):
            continue
        # 基準日 (次の平日・次の日祝) に走る便だけを使う (冬ダイヤなど、期間の違うダイヤが同じデータに入っているため)
        days = [k for k in ('weekday', 'holiday') if run(t['service_id'], ref[k])]
        st = times.get(t['trip_id'], [])
        if not days or len(st) < 2:
            continue
        patterns.setdefault((t['route_id'], tuple(r['stop_id'] for r in st)), []).append((t, st, days))
    out = []
    for k, ((route_id, stop_ids), trips) in enumerate(sorted(patterns.items())):
        route = routes[route_id]
        names = [stops[s]['stop_name'] for s in stop_ids]
        coords = [rnd([float(stops[s]['stop_lon']), float(stops[s]['stop_lat'])]) for s in stop_ids]
        road = gu.road_route(coords)
        vias = gu.split_at_stops(road, coords) if road else [[] for _ in coords[1:]]
        vias = [gu.simplify([a] + v + [b], 20)[1:-1] for a, b, v in zip(coords, coords[1:], vias)]
        rows = [[names[0], coords[0], 1]]
        for n, c, v in zip(names[1:], coords[1:], vias):
            rows.extend(['', p, 0] for p in v)
            rows.append([n, c, 1])
        seen, trip_rows = set(), []
        for t, st, days in trips:
            dep0 = gtfs_sec(st[0]['departure_time'] or st[0]['arrival_time'])
            tt = [[gtfs_sec(r['arrival_time'] or r['departure_time']) - dep0,
                   gtfs_sec(r['departure_time'] or r['arrival_time']) - dep0] for r in st]
            key = (dep0, json.dumps(tt), tuple(days))
            if key in seen:
                continue
            seen.add(key)
            trip_rows.append({'dep': dep0, 't': tt, 'days': days})
        trip_rows.sort(key=lambda x: x['dep'])
        long_name = route.get('route_long_name') or route.get('route_short_name') or f'{names[0]}〜{names[-1]}'
        op = agency.get(route.get('agency_id'), '')
        sid = f'{prefix}{k}'
        out.append({**BUS, 'id': sid, 'name': f'高速バス {long_name}', 'group': 'highway_bus', 'line': sid,
                    'color': HIGHWAY_COLOR, 'loop': False, 'both': False, 'offset': 0, 'path': rows, 'trips': trip_rows,
                    'note': f'運行: {op}' if op else ''})
        HIGHWAY_LINES.append({'id': sid, 'name': f'高速バス {long_name}', 'operator': op, 'group': 'highway_bus', 'kind': 'bus',
                              'color': HIGHWAY_COLOR, 'stations': [[n, c] for n, c in zip(names, coords)],
                              'shape': [r[1] for r in rows]})
    return out


try:
    hoda = dict(gu.hoda_resources())
    hwy_url = next(u for n, u in hoda.items() if n.startswith('高速バス'))
    highway = highway_services(gu.fetch(hwy_url, 'hoda_highway_bus.zip'), 'hwy')
    highway += highway_services(gu.fetch(NEMURO_GTFS_URL, 'nemuro_bus.zip'), 'hwy_nemuro',
                                keep_route=lambda r: r['route_id'] in NEMURO_HIGHWAY_ROUTES)
    services.extend(highway)
    CREDITS.append('北海道オープンデータプラットフォーム 高速バス GTFS')
    CREDITS.append('根室交通 GTFS')
    print(f'  都市間高速バス: {len(highway)} 系統、{sum(len(s["trips"]) for s in highway)} 便')
except (OSError, StopIteration) as e:
    print('  都市間高速バスの時刻表データを取得できないため、高速バスは表示しない:', e)

# JR北海道 普通・快速 (運転間隔はすべて推計)
LOCALS = [
    # id, 名前, 経路, 仕様, 運転間隔の帯
    ('hakodate_liner', 'はこだてライナー', [('jr_hakodate_s', '函館', '新函館北斗')], {**RAIL, 'cars': 3, 'speed': 85},
     [['06:00', '23:00', 40]], None),
    ('hakodate_mori', '函館本線 普通', [('jr_hakodate_s', '函館', '森')], LOCAL1, [['06:00', '21:00', 120]], None),
    ('hakodate_sawara', '函館本線 普通 (砂原支線経由)', [('jr_hakodate_s', '函館', '大沼'), ('jr_hakodate_sawara', '大沼', '森')],
     LOCAL1, [['07:30', '19:00', 240]], None),
    ('hakodate_oshamambe', '函館本線 普通', [('jr_hakodate_s', '森', '長万部')], LOCAL1, [['06:00', '20:00', 180]], None),
    ('yamasen_kutchan', '函館本線 普通 (山線)', [('jr_hakodate_yama', '長万部', '倶知安')], LOCAL1, [['06:00', '19:00', 180]], None),
    ('yamasen_otaru', '函館本線 普通 (山線)', [('jr_hakodate_yama', '倶知安', '小樽')], RAIL, [['06:00', '21:30', 60]], None),
    ('sapporo_otaru', '函館本線 普通', [('jr_hakodate_n', '小樽', '江別')], {**RAIL, 'cars': 6, 'speed': 80},
     [['05:40', '07:00', 30], ['07:00', '09:00', 15], ['09:00', '23:30', 20]], None),
    ('sapporo_iwamizawa', '函館本線 普通', [('jr_hakodate_n', '札幌', '岩見沢')], {**RAIL, 'cars': 4, 'speed': 80},
     [['06:00', '23:30', 30]], None),
    ('iwamizawa_asahikawa', '函館本線 普通', [('jr_hakodate_n', '岩見沢', '旭川')], RAIL, [['06:00', '21:30', 90]], None),
    ('chitose_local', '千歳線 普通', [('jr_chitose', '札幌', '千歳')], {**RAIL, 'cars': 4, 'speed': 80},
     [['06:00', '23:30', 30]], None),
    ('chitose_tomakomai', '千歳線 普通', [('jr_chitose', '千歳', '沼ノ端'), ('jr_muroran_n', '沼ノ端', '苫小牧')], RAIL,
     [['06:00', '22:30', 60]], None),
    ('muroran_local', '室蘭本線 普通', [('jr_muroran', '苫小牧', '東室蘭'), ('jr_muroran_branch', '東室蘭', '室蘭')], RAIL,
     [['06:00', '22:00', 60]], None),
    ('muroran_oshamambe', '室蘭本線 普通', [('jr_muroran', '東室蘭', '長万部')], LOCAL1, [['06:00', '20:00', 180]], None),
    ('muroran_iwamizawa', '室蘭本線 普通', [('jr_muroran_n', '苫小牧', '岩見沢')], LOCAL1, [['06:00', '20:30', 150]], None),
    ('hidaka_local', '日高本線 普通', [('jr_hidaka', '苫小牧', '鵡川')], LOCAL1, [['06:00', '21:00', 120]], None),
    ('sekisho_local', '石勝線 普通', [('jr_sekisho', '南千歳', '新夕張')], LOCAL1, [['06:30', '19:30', 180]], None),
    ('sassho_local', '学園都市線 普通', [('jr_sassho', '札幌', '北海道医療大学')], {**RAIL, 'cars': 4, 'speed': 80},
     [['06:00', '23:30', 30]], None),
    ('sassho_ainosato', '学園都市線 普通', [('jr_sassho', '札幌', 'あいの里公園')], {**RAIL, 'cars': 4, 'speed': 80},
     [['06:15', '22:00', 30]], None),
    ('furano_biei', '富良野線 普通', [('jr_furano', '旭川', '美瑛')], LOCAL1, [['06:00', '22:00', 60]], None),
    ('furano_local', '富良野線 普通', [('jr_furano', '美瑛', '富良野')], LOCAL1, [['06:30', '20:30', 120]], None),
    ('nemuro_takikawa', '根室本線 普通', [('jr_nemuro_w', '滝川', '富良野')], LOCAL1, [['06:00', '20:00', 120]], None),
    ('nemuro_obihiro', '根室本線 普通', [('jr_nemuro', '新得', '池田')], LOCAL1, [['06:00', '21:00', 120]], None),
    ('nemuro_kushiro', '根室本線 普通', [('jr_nemuro', '池田', '釧路')], LOCAL1, [['06:00', '20:00', 180]], None),
    ('hanasaki', '花咲線 普通', [('jr_hanasaki', '釧路', '根室')], LOCAL1, [['06:00', '19:00', 180]], None),
    ('soya_nayoro', '宗谷本線 普通', [('jr_soya', '旭川', '名寄')], LOCAL1, [['06:00', '21:30', 90]], None),
    ('soya_wakkanai', '宗谷本線 普通', [('jr_soya', '名寄', '稚内')], LOCAL1, [['06:00', '18:00', 240]], None),
    ('sekihoku_kamikawa', '石北本線 普通', [('jr_sekihoku', '旭川', '上川')], LOCAL1, [['06:00', '20:00', 150]], None),
    ('sekihoku_kitami', '石北本線 普通', [('jr_sekihoku', '遠軽', '北見')], LOCAL1, [['06:00', '19:00', 180]], None),
    ('sekihoku_abashiri', '石北本線 普通', [('jr_sekihoku', '北見', '網走')], LOCAL1, [['06:00', '21:00', 120]], None),
    ('senmo_shari', '釧網本線 普通', [('jr_senmo', '網走', '知床斜里')], LOCAL1, [['06:00', '20:00', 120]], None),
    ('senmo_kushiro', '釧網本線 普通', [('jr_senmo', '知床斜里', '釧路')], LOCAL1, [['06:00', '18:00', 180]], None),
]
for i, (sid, name, legs, spec, bands, stops) in enumerate(LOCALS):
    service(sid, name, 'jr', legs[0][0], JR_COLOR, legs, spec, stops, bands=bands, offset=(i * 7) % 20)

service('airport', '快速 エアポート', 'jr', 'jr_chitose', '#1E88E5',
        [('jr_chitose', '札幌', '南千歳'), ('jr_chitose_airport', '南千歳', '新千歳空港')],
        {**RAIL, 'cars': 6, 'speed': 110, 'dwell': 40}, ['札幌', '新札幌', '北広島', '恵庭', '千歳', '南千歳', '新千歳空港'],
        bands=[['05:40', '06:30', 20], ['06:30', '23:00', 12]], offset=3)

# 特急 (停車駅は主なもの。本数・時刻は推計)
SAP_TOMAKOMAI = [('jr_chitose', '札幌', '沼ノ端'), ('jr_muroran_n', '沼ノ端', '苫小牧')]
SAP_HAKO = SAP_TOMAKOMAI + [('jr_muroran', '苫小牧', '長万部'), ('jr_hakodate_s', '長万部', '函館')]
LIMITED = [
    ('hokuto', '特急 北斗', SAP_HAKO, '#E53935',
     ['札幌', '新札幌', '南千歳', '苫小牧', '登別', '東室蘭', '伊達紋別', '洞爺', '長万部', '八雲', '森', '大沼公園', '新函館北斗', '五稜郭', '函館'],
     {'bands': [['06:00', '19:30', 75]]}),
    ('suzuran', '特急 すずらん', SAP_TOMAKOMAI + [('jr_muroran', '苫小牧', '東室蘭'), ('jr_muroran_branch', '東室蘭', '室蘭')], '#8E24AA',
     ['札幌', '新札幌', '南千歳', '苫小牧', '白老', '登別', '幌別', '東室蘭', '室蘭'],
     {'bands': [['07:00', '21:00', 150]], 'offset': 20}),
    ('lilac', '特急 ライラック・カムイ', [('jr_hakodate_n', '札幌', '旭川')], '#2E7D32',
     ['札幌', '岩見沢', '美唄', '砂川', '滝川', '深川', '旭川'],
     {'bands': [['06:30', '22:30', 30]], 'speed': 125}),
    ('soya', '特急 宗谷', [('jr_hakodate_n', '札幌', '旭川'), ('jr_soya', '旭川', '稚内')], '#FFB300',
     ['札幌', '岩見沢', '滝川', '深川', '旭川', '和寒', '士別', '名寄', '美深', '音威子府', '幌延', '豊富', '南稚内', '稚内'],
     {'departures': ['07:30'], 'departuresReturn': ['17:45'], 'speed': 85}),
    ('sarobetsu', '特急 サロベツ', [('jr_soya', '旭川', '稚内')], '#FFD54F',
     ['旭川', '和寒', '士別', '名寄', '美深', '音威子府', '幌延', '豊富', '南稚内', '稚内'],
     {'departures': ['12:45', '18:00'], 'departuresReturn': ['06:40', '13:30'], 'speed': 80}),
    ('okhotsk', '特急 オホーツク', [('jr_hakodate_n', '札幌', '旭川'), ('jr_sekihoku', '旭川', '網走')], '#26A69A',
     ['札幌', '岩見沢', '滝川', '深川', '旭川', '上川', '丸瀬布', '遠軽', '生田原', '留辺蘂', '北見', '端野', '美幌', '女満別', '網走'],
     {'departures': ['06:50', '17:30'], 'departuresReturn': ['06:20', '17:25'], 'speed': 78}),
    ('taisetsu', '特急 大雪', [('jr_sekihoku', '旭川', '網走')], '#80CBC4',
     ['旭川', '上川', '丸瀬布', '遠軽', '生田原', '留辺蘂', '北見', '端野', '美幌', '女満別', '網走'],
     {'departures': ['10:40', '18:10'], 'departuresReturn': ['09:30', '13:40'], 'speed': 72}),
    ('oozora', '特急 おおぞら', [('jr_chitose', '札幌', '南千歳'), ('jr_sekisho', '南千歳', '新得'), ('jr_nemuro', '新得', '釧路')], '#1E88E5',
     ['札幌', '新札幌', '南千歳', '追分', 'トマム', '新得', '芽室', '帯広', '池田', '浦幌', '白糠', '釧路'],
     {'bands': [['06:50', '19:30', 150]]}),
    ('tokachi', '特急 とかち', [('jr_chitose', '札幌', '南千歳'), ('jr_sekisho', '南千歳', '新得'), ('jr_nemuro', '新得', '帯広')], '#90CAF9',
     ['札幌', '新札幌', '南千歳', 'トマム', '新得', '十勝清水', '芽室', '帯広'],
     {'bands': [['07:50', '21:00', 180]], 'offset': 30}),
]
for sid, name, legs, color, stops, extra in LIMITED:
    service(sid, name, 'jr_ltd', legs[-1][0], color, legs, LTD, stops, **extra)

service('hayabusa', '北海道新幹線 はやぶさ', 'jr_ltd', 'shinkansen', '#7B5EA7', [('shinkansen', '新函館北斗', '新青森')],
        {'kind': 'rail', 'cars': 10, 'carLength': 25, 'width': 3.4, 'height': 4.0, 'speed': 160, 'dwell': 60, 'accel': 60},
        bands=[['06:35', '21:30', 70]])

# 道南いさりび鉄道
service('isaribi_local', '道南いさりび鉄道 普通', 'hokkaido_other', 'isaribi', '#1565C0', [('isaribi', '木古内', '函館')],
        LOCAL1, bands=[['06:00', '22:00', 90]])
service('isaribi_kamiiso', '道南いさりび鉄道 普通', 'hokkaido_other', 'isaribi', '#1565C0', [('isaribi', '上磯', '函館')],
        LOCAL1, bands=[['06:30', '21:00', 90]], offset=45)

# ------------------------------------------------------------------ フェリー・旅客船 (航路の形は概略、時刻は推計)
FERRY_LINES = []
BIG = {'kind': 'ship', 'cars': 1, 'carLength': 190, 'width': 27, 'height': 22, 'dwell': 1800, 'accel': 300}
MID = {'kind': 'ship', 'cars': 1, 'carLength': 95, 'width': 16, 'height': 13, 'dwell': 600, 'accel': 120}
SMALL = {'kind': 'ship', 'cars': 1, 'carLength': 45, 'width': 9, 'height': 8, 'dwell': 300, 'accel': 60}

HAKODATE_PORT = ('函館港', [140.7065, 41.8010])
AOMORI_PORT = ('青森港', [140.7010, 40.8420])
OMA_PORT = ('大間港', [140.9050, 41.5310])
TOMAKOMAI_PORT = ('苫小牧港', [141.6200, 42.6300])
HACHINOHE_PORT = ('八戸港', [141.5350, 40.5480])
OTARU_PORT = ('小樽港', [141.0210, 43.1970])
WAKKANAI_PORT = ('稚内港', [141.6880, 45.4140])
OSHIDOMARI = ('鴛泊港（利尻島）', [141.2430, 45.2430])
KAFUKA = ('香深港（礼文島）', [141.0450, 45.3040])
HABORO = ('羽幌港', [141.6950, 44.3650])
YAGISHIRI = ('焼尻港', [141.4250, 44.4350])
TEURI = ('天売港', [141.3350, 44.4150])
ESASHI = ('江差港', [140.1250, 41.8680])
OKUSHIRI = ('奥尻港', [139.5200, 42.1700])


def duration(coords, stops, spec, speed):
    """sim.js と同じ計算で、始発から終点までの所要時間 [秒]"""
    v = speed / 3.6
    t, last, first = 0.0, 0, True
    cum = [0.0]
    for i in range(1, len(coords)):
        cum.append(cum[-1] + haversine(coords[i - 1], coords[i]))
    for i in range(1, len(coords)):
        if not (stops[i] or i == len(coords) - 1):
            continue
        t += (cum[i] - cum[last]) / v + spec['accel'] * 2
        if i != len(coords) - 1:
            t += spec['dwell']
        last = i
    return t


def hhmm(sec):
    sec = int(round(sec)) % 86400
    return f'{sec // 3600:02d}:{sec // 60 % 60:02d}'


def sec_of(s):
    h, m = map(int, s.split(':'))
    return h * 3600 + m * 60


def ferry(sid, name, points, spec, minutes, departures=None, departures_return=None, arrivals_return=None, bands=None, note='運航パターンは推計。航路の形は概略'):
    """points: 港 (名前, 座標) と途中の点 ([lon, lat]) の並び。minutes: 始発港から終点までの所要時間 [分]"""
    raw = [(p[0], p[1], 1) if isinstance(p, tuple) else ('', p, 0) for p in points]
    coords = [r[1] for r in raw]
    stops = [r[2] for r in raw]
    # 所要時間に合わせて速度を決める
    lo, hi = 1.0, 200.0
    for _ in range(60):
        mid = (lo + hi) / 2
        if duration(coords, stops, spec, mid) > minutes * 60:
            lo = mid
        else:
            hi = mid
    speed = round(hi, 1)
    path = []
    dense = densify(coords, 2000)
    # 港は densify 後も元の座標のまま残る (始点・終点・各区間の端点)
    ports = {tuple(c): n for n, c, s in raw if s}
    for c in dense:
        n = ports.get(tuple(c), '')
        path.append([n, rnd(c), 1 if n else 0])
    sv = {**spec, 'id': sid, 'name': name, 'group': 'ferry', 'line': sid, 'color': '#e8eef5', 'speed': speed,
          'loop': False, 'both': True, 'offset': 0, 'note': note, 'path': path}
    if bands:
        sv['bands'] = bands
    if departures:
        sv['departures'] = departures
    if arrivals_return:
        # 帰りは終点 (沖合) を出る時刻を、始発港に着く時刻から逆算する
        d = duration(list(reversed(coords)), list(reversed(stops)), spec, speed)
        departures_return = [hhmm(sec_of(a) - d) for a in arrivals_return]
    if departures_return:
        sv['departuresReturn'] = sorted(departures_return)
    services.append(sv)
    FERRY_LINES.append({'id': sid, 'name': name, 'operator': name.split('（')[0], 'group': 'ferry', 'kind': 'ship', 'color': '#e8eef5',
                        'stations': [[n, c] for n, c, s in raw if s], 'shape': [rnd(c) for c in coords]})


ferry('ferry_seikan', '津軽海峡フェリー・青函フェリー（函館〜青森）',
      [HAKODATE_PORT, [140.690, 41.770], [140.665, 41.720], [140.660, 41.600], [140.690, 41.440], [140.725, 41.250],
       [140.740, 41.120], [140.750, 40.970], [140.725, 40.875], AOMORI_PORT],
      BIG, 225, bands=[['00:30', '23:59', 90]])
ferry('ferry_oma', '津軽海峡フェリー（函館〜大間）',
      [HAKODATE_PORT, [140.690, 41.770], [140.665, 41.720], [140.750, 41.620], [140.880, 41.550], OMA_PORT],
      MID, 90, departures=['09:30', '16:00'], departures_return=['07:00', '13:30'])
ferry('ferry_hachinohe', 'シルバーフェリー（苫小牧〜八戸）',
      [TOMAKOMAI_PORT, [141.630, 42.550], [141.800, 42.000], [141.850, 41.500], [141.700, 40.900], [141.570, 40.580], HACHINOHE_PORT],
      BIG, 450, departures=['01:00', '05:00', '13:30', '21:15'], departures_return=['00:30', '08:30', '17:00', '22:00'])
ferry('ferry_pacific', 'さんふらわあ・太平洋フェリー（苫小牧〜大洗・仙台）',
      [TOMAKOMAI_PORT, [141.630, 42.550], [141.900, 42.100], [142.100, 41.600], ('大洗・仙台方面', [142.150, 41.300])],
      BIG, 240, departures=['01:30', '18:45', '19:00'], arrivals_return=['11:00', '13:30', '19:45'],
      note='運航パターンは推計。航路の形は概略で、北海道の沖合で出現・消滅します')
ferry('ferry_shinnihonkai', '新日本海フェリー（小樽〜新潟・舞鶴）',
      [OTARU_PORT, [141.050, 43.260], [140.700, 43.400], [140.350, 43.450], [140.100, 43.300], ('新潟・舞鶴方面', [139.950, 43.000])],
      BIG, 240, departures=['17:00', '23:30'], arrivals_return=['04:30', '20:45'],
      note='運航パターンは推計。航路の形は概略で、北海道の沖合で出現・消滅します')
ferry('ferry_rishiri', 'ハートランドフェリー（稚内〜利尻島）',
      [WAKKANAI_PORT, [141.700, 45.440], [141.660, 45.475], [141.600, 45.465], [141.400, 45.330], OSHIDOMARI],
      MID, 100, departures=['07:15', '11:25', '15:30'], departures_return=['09:25', '13:40', '17:30'])
ferry('ferry_rebun', 'ハートランドフェリー（稚内〜礼文島）',
      [WAKKANAI_PORT, [141.700, 45.440], [141.660, 45.475], [141.600, 45.465], [141.350, 45.400], [141.100, 45.330], KAFUKA],
      MID, 115, departures=['06:20', '10:35', '14:40'], departures_return=['08:40', '12:50', '16:45'])
ferry('ferry_rishiri_rebun', 'ハートランドフェリー（利尻島〜礼文島）',
      [OSHIDOMARI, [141.180, 45.290], KAFUKA],
      MID, 45, departures=['09:10', '14:20'], departures_return=['08:10', '15:30'])
ferry('ferry_haboro', '羽幌沿海フェリー（羽幌〜焼尻島〜天売島）',
      [HABORO, YAGISHIRI, TEURI],
      SMALL, 75, departures=['08:30', '13:30'], departures_return=['10:30', '15:40'])
ferry('ferry_okushiri', 'ハートランドフェリー（江差〜奥尻島）',
      [ESASHI, OKUSHIRI],
      MID, 130, departures=['12:30'], departures_return=['09:00'])

# ------------------------------------------------------------------ 空港・航空便 (便数・時刻は推計)
# 空港の位置・敷地・滑走路・運用時間は「国土数値情報（空港データ C28、2021 年度）」、
# 利用状況は国土交通省「空港管理状況調書」(令和 7 年、暦年・年度別) から取る。取得できなければ下の概略値を使う
C28_URL = 'https://nlftp.mlit.go.jp/ksj/gml/data/C28/C28-21/C28-21_GML.zip'
KANRI_URL = 'https://www.mlit.go.jp/koku/content/002016480.xlsx'        # 令和 7 年空港管理状況調書 (月別)
KANRI_TREND_URL = 'https://www.mlit.go.jp/koku/content/002018023.xlsx'  # 暦年・年度別空港管理状況調書 (H28〜R7)
STATS_YEAR = 2025

AIRPORTS = [
    # id, 名前, 座標, 滑走路の向き (真方位), 長さ, 運用時間, 国土数値情報・空港管理状況調書での名前
    ('CTS', '新千歳空港', [141.6923, 42.7752], 1, 3000, ['0000', '2400'], '新千歳空港', '新千歳'),
    ('OKD', '丘珠空港', [141.3814, 43.1176], 131, 1500, ['0800', '2000'], '札幌飛行場', '札幌'),
    ('HKD', '函館空港', [140.8219, 41.7700], 111, 3000, ['0730', '2130'], '函館空港', '函館'),
    ('AKJ', '旭川空港', [142.4475, 43.6708], 151, 2500, ['0800', '2130'], '旭川空港', '旭川'),
    ('KUH', 'たんちょう釧路空港', [144.1929, 43.0410], 161, 2500, ['0800', '2130'], '釧路空港', '釧路'),
    ('OBO', 'とかち帯広空港', [143.2172, 42.7333], 161, 2500, ['0800', '2130'], '帯広空港', '帯広'),
    ('MMB', '女満別空港', [144.1642, 43.8806], 171, 2500, ['0800', '2130'], '女満別空港', '女満別'),
    ('SHB', '中標津空港', [144.9600, 43.5775], 71, 2000, ['0830', '2030'], '中標津空港', '中標津'),
    ('MBE', 'オホーツク紋別空港', [143.4042, 44.3039], 131, 2000, ['0830', '2030'], '紋別空港', '紋別'),
    ('WKJ', '稚内空港', [141.8008, 45.4042], 71, 2200, ['0830', '2030'], '稚内空港', '稚内'),
    ('RIS', '利尻空港', [141.1864, 45.2420], 61, 1800, ['0830', '1830'], '利尻空港', '利尻'),
    ('OIR', '奥尻空港', [139.4331, 42.0717], 131, 1500, ['0830', '1830'], '奥尻空港', '奥尻'),
]


def fetch(url, name):
    path = os.path.join(CACHE, name)
    if not os.path.exists(path):
        print('download', url)
        urllib.request.urlretrieve(url, path)
    return path


def polygon_area(ring):
    """おおよその面積 [m²] (平面近似)"""
    lat0 = ring[0][1]
    kx, ky = 111320 * math.cos(math.radians(lat0)), 110540
    s = 0.0
    for (x1, y1), (x2, y2) in zip(ring, ring[1:]):
        s += (x1 * kx) * (y2 * ky) - (x2 * kx) * (y1 * ky)
    return abs(s) / 2


def long_axis(ring):
    """敷地の形の長軸の向き (真方位 0〜180 度)。滑走路の向きとみなす"""
    pts = densify(ring, 50)
    lat0 = sum(p[1] for p in pts) / len(pts)
    lon0 = sum(p[0] for p in pts) / len(pts)
    kx, ky = 111320 * math.cos(math.radians(lat0)), 110540
    xs = [(p[0] - lon0) * kx for p in pts]
    ys = [(p[1] - lat0) * ky for p in pts]
    sxx = sum(x * x for x in xs)
    syy = sum(y * y for y in ys)
    sxy = sum(x * y for x, y in zip(xs, ys))
    theta = 0.5 * math.atan2(2 * sxy, sxx - syy)  # x 軸 (東) からの角度
    return round((90 - math.degrees(theta)) % 180, 1)


def load_c28():
    """C28 の空港名 → 位置 (標点)・敷地・滑走路の長さ・運用時間"""
    zpath = fetch(C28_URL, 'C28-21_GML.zip')
    with zipfile.ZipFile(zpath) as z:
        def geo(k):
            with z.open(f'UTF-8/C28-21_{k}.geojson') as f:
                return json.load(f)['features']
        ref = {f['properties']['C28_000']: f['geometry']['coordinates'] for f in geo('AirportReferencePoint')}
        out = {}
        for f in geo('Airport'):
            p = f['properties']
            ring = f['geometry']['coordinates'][0]
            cur = out.get(p['C28_005'])
            # 敷地が複数に分かれている空港 (例: 稚内) は、いちばん大きいものを使う
            if cur and polygon_area(cur['polygon']) >= polygon_area(ring):
                continue
            out[p['C28_005']] = {
                'coord': [round(x, 6) for x in ref[p['C28_101'].lstrip('#')]],
                'polygon': [rnd(c) for c in ring],
                'runway': max(int(x) for x in p['C28_012'].split(',')),
                'hours': [p['C28_009'], p['C28_010']],
            }
    return out


def load_kanri():
    """空港管理状況調書の空港名 → 利用状況 (年間・月別・10 年の推移)"""
    import openpyxl

    def sheet(path):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        return list(wb[wb.sheetnames[-1]].iter_rows(values_only=True))

    def blocks(rows):
        """「空港名 ：」の行ごとに区切る"""
        out, name = {}, None
        for r in rows:
            for j, c in enumerate(r):
                if isinstance(c, str) and c.replace(' ', '').replace('　', '').startswith('空港名'):
                    raw = next(x for x in r[j + 1:] if x)
                    name = re.sub(r'（.*?）', '', raw).replace(' ', '').replace('　', '')
                    out[name] = []
                    break
            else:
                if name:
                    out[name].append(r)
        return out

    monthly = blocks(sheet(fetch(KANRI_URL, 'kanri_r7.xlsx')))
    trend = blocks(sheet(fetch(KANRI_TREND_URL, 'kanri_trend.xlsx')))
    stats = {}
    for name, rows in monthly.items():
        # 列: 月・着陸 (国際・国内・計)・乗降客 (国際 乗/降/通過/計・国内 乗/降/計)・合計
        mrows, total = [], None
        for r in rows:
            label = r[2] if len(r) > 2 else None
            if isinstance(label, str) and re.fullmatch(r'\d{1,2}月', label) and total is None and len(mrows) < 12:
                mrows.append((int(label[:-1]), r))
            elif isinstance(label, str) and label.replace(' ', '') == '暦年計':
                total = r
                break
        if total is None or len(mrows) != 12:
            continue
        num = lambda v: int(v or 0)
        st = {'year': STATS_YEAR, 'passengers': num(total[13]), 'landings': num(total[5]),
              'domestic': num(total[12]), 'international': num(total[9]),
              'monthly': [{'month': m, 'passengers': num(r[13]), 'landings': num(r[5])} for m, r in mrows], 'trend': []}
        # 暦年・年度別: 年 (H28〜R7) の行が暦年、続いて年度の行が並ぶ。暦年の 10 行を使う
        era = {'28': 2016, '29': 2017, '30': 2018, '元': 2019}
        for r in trend.get(name, []):
            y = r[1] if len(r) > 1 else None
            if y is None:
                continue
            y = str(y).strip()
            year = era.get(y) or (2018 + int(y) if y.isdigit() and int(y) <= 20 else None)
            if year and not any(t['year'] == year for t in st['trend']) and isinstance(r[12], (int, float)):
                st['trend'].append({'year': year, 'passengers': int(r[12]), 'landings': int(r[4] or 0)})
        st['trend'].sort(key=lambda t: t['year'])
        stats[name] = st
    return stats


AIRPORT_INFO, AIRPORT_STATS = {}, {}
try:
    C28 = load_c28()
    for ap in AIRPORTS:
        code, name, c, hdg, runway, hours, c28_name, _ = ap
        d = C28.get(c28_name)
        if d is None:
            print(f'  {name}: 国土数値情報（空港データ）に無いため概略値を使う')
            continue
        axis = long_axis(d['polygon'])
        # 滑走路の向きは敷地の長軸。離着陸の向きは、概略値に近いほう (axis か axis+180) にする
        diff = lambda a, b: abs((a - b + 180) % 360 - 180)
        heading = axis if diff(axis, hdg) <= diff(axis + 180, hdg) else (axis + 180) % 360
        # 敷地の形が滑走路と揃っていない空港 (新千歳・丘珠) は、長軸が滑走路からずれるので概略値を使う
        if diff(heading, hdg) > 10:
            print(f'  {name}: 敷地の長軸 {heading:.0f}° が滑走路の向き {hdg}° と {diff(heading, hdg):.0f}° ずれるため、滑走路の向きは概略値を使う')
            heading = hdg
        AIRPORT_INFO[code] = {**d, 'heading': round(heading, 1)}
except OSError as e:
    print('  国土数値情報（空港データ）を取得できないため、空港の位置は概略値を使う:', e)
try:
    KANRI = load_kanri()
    for ap in AIRPORTS:
        if ap[7] in KANRI:
            AIRPORT_STATS[ap[0]] = KANRI[ap[7]]
        else:
            print(f'  {ap[1]}: 空港管理状況調書に無いため、利用状況は表示しない')
except (OSError, ImportError) as e:
    print('  空港管理状況調書を読めないため、利用状況は表示しない:', e)

# 以降の航空便の計算は、国土数値情報の位置・向き・運用時間を使う
AIRPORTS = [(code, name,
             AIRPORT_INFO.get(code, {}).get('coord', c),
             AIRPORT_INFO.get(code, {}).get('heading', hdg),
             AIRPORT_INFO.get(code, {}).get('runway', runway),
             AIRPORT_INFO.get(code, {}).get('hours', hours))
            for code, name, c, hdg, runway, hours, _, _ in AIRPORTS]
AP = {a[0]: a for a in AIRPORTS}

DEST = {
    '羽田': [139.7798, 35.5494], '成田': [140.3864, 35.7647], '伊丹': [135.4383, 34.7855], '関西': [135.2440, 34.4347],
    '中部': [136.8053, 34.8584], '神戸': [135.2239, 34.6328], '福岡': [130.4511, 33.5859], '仙台': [140.9170, 38.1397],
    '那覇': [127.6461, 26.1958], '新潟': [139.1206, 37.9559], '広島': [132.9194, 34.4361], '茨城': [140.4147, 36.1811],
    '花巻': [141.1353, 39.4286], '三沢': [141.3683, 40.7033], '松本': [137.9228, 36.1667],
    'ソウル（仁川）': [126.4505, 37.4602], '台北（桃園）': [121.2328, 25.0797], '香港': [113.9185, 22.3080],
    '上海（浦東）': [121.8053, 31.1443], 'バンコク': [100.7501, 13.6900], 'シンガポール': [103.9915, 1.3644],
}
INTL = {'ソウル（仁川）', '台北（桃園）', '香港', '上海（浦東）', 'バンコク', 'シンガポール'}

# 道外との便: 空港ごとに (行き先, 1 日の便数)
EXTERNAL = {
    'CTS': [('羽田', 48), ('成田', 14), ('伊丹', 12), ('関西', 9), ('中部', 11), ('神戸', 5), ('福岡', 8), ('仙台', 7),
            ('那覇', 3), ('新潟', 3), ('広島', 2), ('茨城', 2), ('花巻', 2),
            ('ソウル（仁川）', 10), ('台北（桃園）', 6), ('香港', 3), ('上海（浦東）', 2), ('バンコク', 2), ('シンガポール', 1)],
    'OKD': [('三沢', 1), ('松本', 1)],
    'HKD': [('羽田', 8), ('伊丹', 2), ('中部', 1), ('成田', 1), ('台北（桃園）', 1)],
    'AKJ': [('羽田', 7), ('中部', 1)],
    'KUH': [('羽田', 6)],
    'OBO': [('羽田', 7)],
    'MMB': [('羽田', 6), ('中部', 1)],
    'SHB': [('羽田', 1)],
    'MBE': [('羽田', 1)],
    'WKJ': [('羽田', 1)],
}
# 道内の便: (空港, 空港, 片道 1 日の便数)
INTERNAL = [
    ('OKD', 'HKD', 3), ('OKD', 'KUH', 3), ('OKD', 'MMB', 2), ('OKD', 'RIS', 1), ('OKD', 'OIR', 1),
    ('CTS', 'MMB', 3), ('CTS', 'WKJ', 2), ('CTS', 'RIS', 1), ('CTS', 'SHB', 1), ('HKD', 'OIR', 1),
]

PLANE = {'kind': 'plane', 'cars': 1, 'carLength': 40, 'width': 36, 'height': 6, 'speed': 480, 'dwell': 0, 'accel': 150}
CRUISE = 6000
CLIMB = 80      # 上昇 [m/km]
DESCENT = 52.5  # 降下 [m/km] (約 3 度)
DEP_LEN = 170_000
ARR_LEN = 200_000


def threshold(ap):
    _, _, c, hdg, runway, _ = ap
    return offset(c, hdg + 180, runway / 2)


def dep_track(ap, target, length):
    """離陸: 滑走路の向きにまっすぐ上昇し、行き先の方向へ旋回する (1km ごとの点)"""
    hdg = ap[3]
    pts = [threshold(ap)]
    brg = hdg
    for k in range(1, int(length // 1000) + 1):
        if k > 8:
            want = bearing(pts[-1], target)
            diff = (want - brg + 540) % 360 - 180
            brg += max(-4, min(4, diff))
        pts.append(offset(pts[-1], brg, 1000))
    return pts


def arr_track(ap, origin, length):
    """着陸: 行き先 (出発地) の方向から回り込み、滑走路の向きに合わせて降りる"""
    hdg = ap[3]
    back = [offset(threshold(ap), hdg, 1500), threshold(ap)]  # 着地後 1.5km 滑走
    brg = hdg + 180
    for k in range(1, int(length // 1000) + 1):
        if k > 12:
            want = bearing(back[-1], origin)
            diff = (want - brg + 540) % 360 - 180
            brg += max(-4, min(4, diff))
        back.append(offset(back[-1], brg, 1000))
    return list(reversed(back))


def alt_profile(coords, climb=True, descend=True):
    cum = [0.0]
    for i in range(1, len(coords)):
        cum.append(cum[-1] + haversine(coords[i - 1], coords[i]))
    total = cum[-1]
    out = []
    for d in cum:
        a = CRUISE
        if climb:
            a = min(a, max(0.0, (d - 1000) / 1000 * CLIMB))
        if descend:
            a = min(a, max(0.0, (total - d - 1500) / 1000 * DESCENT))
        out.append([round(d), round(a)])
    return out


def spread(n, start, end, shift):
    """n 便を start〜end [分] に均等に並べる"""
    span = end - start
    return [hhmm((start + span * (i + 0.5) / n + shift) * 60) for i in range(n)]


def hours_min(ap):
    a, b = ap[5]
    s = int(a[:2]) * 60 + int(a[2:])
    e = int(b[:2]) * 60 + int(b[2:])
    if e - s >= 1440:
        s, e = 7 * 60, 21 * 60 + 30
    return s, e


def path_rows(coords, first, last):
    rows = [['', rnd(c), 0] for c in coords]
    rows[0] = [first, rows[0][1], 1]
    rows[-1] = [last, rows[-1][1], 1]
    return rows


plane_services = []
for code, routes in EXTERNAL.items():
    ap = AP[code]
    s, e = hours_min(ap)
    for i, (dest, n) in enumerate(routes):
        target = DEST[dest]
        dep = dep_track(ap, target, DEP_LEN)
        arr = arr_track(ap, target, ARR_LEN)
        base = {**PLANE, 'group': 'air', 'line': f'air_{code}', 'color': '#f5f7fa', 'loop': False, 'both': False, 'offset': 0,
                'airport': code, 'intl': dest in INTL}
        shift = (i * 11) % 30 - 15
        plane_services.append({**base, 'id': f'air_{code}_{i}_dep', 'name': f'{ap[1]} ⇄ {dest}',
                               'path': path_rows(dep, ap[1], dest), 'alt': alt_profile(dep, descend=False),
                               'departures': spread(n, s + 30, e - 45, shift), 'flight': 'dep',
                               'fromName': ap[1], 'toName': dest})
        # 到着便: 着陸の時刻が運用時間に収まるよう、進入開始の時刻を決める
        arr_dur = path_length(arr) / (PLANE['speed'] / 3.6) + PLANE['accel'] * 2
        lands = spread(n, s + 45, e - 15, -shift)
        plane_services.append({**base, 'id': f'air_{code}_{i}_arr', 'name': f'{dest} ⇄ {ap[1]}',
                               'path': path_rows(arr, dest, ap[1]), 'alt': alt_profile(arr, climb=False),
                               'departures': sorted(hhmm(sec_of(t) - arr_dur) for t in lands), 'flight': 'arr',
                               'fromName': dest, 'toName': ap[1]})

for i, (a, b, n) in enumerate(INTERNAL):
    for j, (x, y) in enumerate(((a, b), (b, a))):
        apx, apy = AP[x], AP[y]
        dist = haversine(apx[2], apy[2])
        # 出発側は旋回して相手の空港へ向かい、到着側は相手の滑走路へ回り込む。中間は直線でつなぐ
        dep = dep_track(apx, apy[2], min(40_000, dist * 0.3))
        arr = arr_track(apy, apx[2], min(60_000, dist * 0.4))
        mid = densify([dep[-1], arr[0]], 1000)[1:-1]
        coords = dep + mid + arr
        s, e = hours_min(apx)
        s2, e2 = hours_min(apy)
        dur = path_length(coords) / (PLANE['speed'] / 3.6) + PLANE['accel'] * 2
        lo, hi = max(s + 30, s2 + 15 - dur / 60), min(e - 45, e2 - 15 - dur / 60)
        plane_services.append({**PLANE, 'id': f'air_in_{i}_{j}', 'name': f'{apx[1]} → {apy[1]}', 'group': 'air',
                               'line': f'air_{x}', 'color': '#f5f7fa', 'loop': False, 'both': False, 'offset': 0,
                               'airport': x, 'intl': False, 'path': path_rows(coords, apx[1], apy[1]),
                               'alt': alt_profile(coords), 'departures': spread(n, lo, hi, (i * 13 + j * 29) % 20 - 10),
                               'flight': 'domestic', 'fromName': apx[1], 'toName': apy[1]})

services.extend(plane_services)


# 滑走路の番号 (両端) と、平行滑走路 (L / R) があるか。実際の飛行機 (ADS-B) の使用滑走路の推定に使う
RUNWAY_NAMES = {
    'CTS': ('01', '19', True), 'OKD': ('14', '32', False), 'HKD': ('12', '30', False), 'AKJ': ('16', '34', False),
    'KUH': ('17', '35', False), 'OBO': ('17', '35', False), 'MMB': ('18', '36', False), 'SHB': ('08', '26', False),
    'MBE': ('14', '32', False), 'WKJ': ('08', '26', False), 'RIS': ('07', '25', False), 'OIR': ('13', '31', False),
}


def runway_polygon(ap):
    _, _, c, hdg, runway, _ = ap
    half_l, half_w = runway / 2 + 200, 150
    a = offset(c, hdg, half_l)
    b = offset(c, hdg + 180, half_l)
    ring = [offset(a, hdg - 90, half_w), offset(a, hdg + 90, half_w), offset(b, hdg + 90, half_w), offset(b, hdg - 90, half_w)]
    ring.append(ring[0])
    return [rnd(p) for p in ring]


airports = []
for ap in AIRPORTS:
    code, name, c, hdg, runway, hours = ap
    routes = [{'dest': d, 'perDay': n, 'intl': d in INTL} for d, n in EXTERNAL.get(code, [])]
    for a, b, n in INTERNAL:
        if code in (a, b):
            routes.append({'dest': AP[b if code == a else a][1].replace('空港', ''), 'perDay': n, 'intl': False})
    # 滑走路の両端の名前 (磁方位の 10 分の 1) と真方位。北海道の磁気偏角は西へ約 9 度
    n1, n2, parallel = RUNWAY_NAMES[code]
    diff = lambda a, b: abs((a - b + 180) % 360 - 180)
    h1 = hdg if diff(hdg, int(n1) * 10 - 9) <= diff(hdg + 180, int(n1) * 10 - 9) else (hdg + 180) % 360
    ends = [{'name': n1, 'heading': round(h1, 1)}, {'name': n2, 'heading': round((h1 + 180) % 360, 1)}]
    airports.append({'id': code, 'name': name, 'coord': c, 'heading': hdg, 'runway': runway, 'hours': hours,
                     'polygon': AIRPORT_INFO.get(code, {}).get('polygon') or runway_polygon(ap),
                     'runwayEnds': ends, 'parallel': parallel,
                     'stats': AIRPORT_STATS.get(code), 'routes': routes})

# ------------------------------------------------------------------ 出力
lines_out = [{'id': lid, 'name': name, 'operator': op, 'group': grp, 'kind': kind, 'color': color,
              'stations': [[n, c] for n, c in sts], 'shape': shape_of(lid)} for lid, name, op, grp, kind, color, sts in LINES]
lines_out += FERRY_LINES
lines_out += HIGHWAY_LINES

with open(os.path.join(HERE, 'holidays.json'), encoding='utf-8') as f:
    holidays = json.load(f)

network = {
    'source': '駅の位置: 国土数値情報（鉄道データ）、駅の並び: 駅データ.jp',
    'groups': GROUPS,
    'lines': lines_out,
    'services': services,
    'airports': airports,
    'calendar': {'holidays': holidays},
    'credits': CREDITS,
    'trackSource': '国土数値情報（鉄道データ）',
}

out = os.path.join(ROOT, 'data', 'network.js')
with open(out, 'w', encoding='utf-8') as f:
    f.write('// 路線・駅・運行パターンのデータ (tools/build_network.py で自動生成)\n')
    f.write('window.NETWORK = ')
    json.dump(network, f, ensure_ascii=False, separators=(',', ':'))
    f.write(';\n')

n_st = len({(n, tuple(c)) for l in lines_out if l['kind'] != 'ship' for n, c in l['stations']})
print(f'{out}: lines {len(lines_out)}, services {len(services)} (planes {len(plane_services)}), stations {n_st}, '
      f'{os.path.getsize(out) / 1e6:.2f} MB')
